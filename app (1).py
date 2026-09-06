import json
import os
import re
from typing import Any

import streamlit as st
from google import genai
from pypdf import PdfReader
from docx import Document


MODEL_NAME = "gemini-2.5-flash"


def get_api_key() -> str:
    """Read the Gemini API key from Streamlit secrets or an environment variable."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return str(st.secrets["GEMINI_API_KEY"])
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY", "").strip()


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a PDF resume."""
    from io import BytesIO

    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table text from a DOCX resume."""
    from io import BytesIO

    doc = Document(BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts).strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract resume text from PDF, DOCX, or TXT."""
    file_bytes = uploaded_file.getvalue()
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]

    if extension == "pdf":
        return extract_pdf_text(file_bytes)
    if extension == "docx":
        return extract_docx_text(file_bytes)
    if extension == "txt":
        return file_bytes.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


def clean_json_response(text: str) -> dict[str, Any]:
    """Convert Gemini's JSON response into a Python dictionary."""
    text = text.strip()

    # Remove Markdown code fences if Gemini returns them.
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text).strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        # Try to recover the first JSON object from surrounding text.
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def analyze_resume(resume_text: str, job_description: str) -> dict[str, Any]:
    """Ask Gemini to evaluate the resume and return structured JSON."""
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "Gemini API key not found. Add GEMINI_API_KEY to Streamlit secrets "
            "or your environment variables."
        )

    client = genai.Client(api_key=api_key)

    jd_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Evaluate against general ATS-friendly resume best practices."
    )

    prompt = f"""
You are an expert ATS resume evaluator and technical recruiter.

Analyze the resume below. Give a practical ATS score from 0 to 100.
If a job description is provided, evaluate keyword and requirement alignment against it.
If no job description is provided, use general ATS-friendly resume standards.

Important:
- Do NOT invent experience, skills, education, employers, metrics, or achievements.
- Distinguish between formatting/readability issues and content gaps.
- Give actionable improvements that the candidate can actually make.
- Treat the score as an estimate, not as a score produced by a specific ATS vendor.
- Keep the response concise but useful.

Return ONLY valid JSON with this exact top-level structure:
{{
  "ats_score": 0,
  "score_label": "Excellent",
  "summary": "Short overall assessment",
  "score_breakdown": {{
    "keyword_alignment": 0,
    "experience_relevance": 0,
    "skills": 0,
    "formatting_and_parseability": 0,
    "impact_and_achievements": 0
  }},
  "strengths": ["...", "..."],
  "critical_issues": ["...", "..."],
  "missing_or_weak_keywords": ["...", "..."],
  "improvements": [
    {{
      "priority": "High",
      "section": "Summary",
      "issue": "What is weak",
      "recommendation": "What to change",
      "example": "A safe example/template based only on information already present"
    }}
  ],
  "ats_checklist": {{
    "contact_information": "Pass",
    "standard_section_headings": "Pass",
    "keyword_coverage": "Needs improvement",
    "quantified_achievements": "Needs improvement",
    "simple_parseable_format": "Pass"
  }}
}}

Scoring guidance:
- keyword_alignment: 0-100
- experience_relevance: 0-100
- skills: 0-100
- formatting_and_parseability: 0-100
- impact_and_achievements: 0-100
- ats_score must be a weighted overall score, not merely the average.
- Use "Excellent", "Strong", "Fair", or "Needs improvement" for score_label.

JOB DESCRIPTION:
{jd_context}

RESUME:
{resume_text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config={
            "temperature": 0.2,
            "response_mime_type": "application/json",
        },
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    return clean_json_response(response.text)


def safe_score(value: Any) -> int:
    try:
        return max(0, min(100, int(float(value))))
    except (TypeError, ValueError):
        return 0


def show_score(score: int, label: str) -> None:
    st.metric("Estimated ATS Score", f"{score}/100", label)
    st.progress(score / 100)


def main() -> None:
    st.set_page_config(
        page_title="Resume ATS Analyzer",
        page_icon="📄",
        layout="wide",
    )

    st.title("📄 Resume ATS Analyzer")
    st.caption(
        "Upload your resume to get an AI-estimated ATS score, weaknesses, "
        "keyword gaps, and practical improvement suggestions."
    )

    with st.sidebar:
        st.header("How it works")
        st.write(
            "1. Upload a PDF, DOCX, or TXT resume.\n"
            "2. Optionally paste a job description.\n"
            "3. Click Analyze Resume.\n"
            "4. Review your score and improvements."
        )
        st.info(
            "Privacy note: the resume text is sent to Gemini for analysis. "
            "Do not upload documents containing information you do not want "
            "processed by the AI service."
        )

    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"],
        help="PDF, DOCX, and TXT are supported.",
        max_upload_size=10,
    )

    job_description = st.text_area(
        "Job description (optional, but recommended)",
        height=220,
        placeholder="Paste the job description here for a more targeted ATS analysis...",
    )

    analyze_clicked = st.button(
        "Analyze Resume",
        type="primary",
        use_container_width=True,
        disabled=uploaded_file is None,
    )

    if not analyze_clicked:
        st.markdown(
            "### What you will get\n"
            "- **ATS score:** estimated score out of 100\n"
            "- **Score breakdown:** keywords, skills, experience, formatting, and impact\n"
            "- **Strengths:** what is already working\n"
            "- **Critical issues:** highest-priority problems\n"
            "- **Keyword gaps:** missing or weak terms\n"
            "- **Action plan:** specific resume improvements"
        )
        return

    if uploaded_file is None:
        st.warning("Please upload a resume first.")
        return

    try:
        with st.spinner("Reading your resume and analyzing it with Gemini..."):
            resume_text = extract_resume_text(uploaded_file)

            if len(resume_text.strip()) < 100:
                st.error(
                    "Very little text could be extracted from this file. "
                    "If this is a scanned/image-only PDF, use a text-based PDF "
                    "or DOCX version of the resume."
                )
                return

            # Prevent accidentally sending an extremely large document.
            resume_text = resume_text[:50000]
            result = analyze_resume(resume_text, job_description)

        score = safe_score(result.get("ats_score"))
        label = str(result.get("score_label", "Assessment"))

        st.success("Analysis complete.")

        show_score(score, label)

        st.subheader("Overall assessment")
        st.write(result.get("summary", "No summary returned."))

        st.subheader("Score breakdown")
        breakdown = result.get("score_breakdown", {})
        cols = st.columns(5)
        breakdown_items = [
            ("Keywords", "keyword_alignment"),
            ("Experience", "experience_relevance"),
            ("Skills", "skills"),
            ("Formatting", "formatting_and_parseability"),
            ("Impact", "impact_and_achievements"),
        ]
        for col, (title, key) in zip(cols, breakdown_items):
            col.metric(title, f"{safe_score(breakdown.get(key))}/100")

        left, right = st.columns(2)

        with left:
            st.subheader("✅ Strengths")
            for item in result.get("strengths", []):
                st.write(f"- {item}")

            st.subheader("⚠️ Critical issues")
            for item in result.get("critical_issues", []):
                st.write(f"- {item}")

        with right:
            st.subheader("🔎 Missing / weak keywords")
            keywords = result.get("missing_or_weak_keywords", [])
            if keywords:
                st.write(", ".join(str(k) for k in keywords))
            else:
                st.write("No major keyword gaps identified.")

            st.subheader("ATS checklist")
            checklist = result.get("ats_checklist", {})
            for key, value in checklist.items():
                readable = key.replace("_", " ").title()
                st.write(f"**{readable}:** {value}")

        st.subheader("🛠️ Recommended improvements")
        improvements = result.get("improvements", [])

        if not improvements:
            st.write("No improvement items were returned.")
        else:
            for index, item in enumerate(improvements, start=1):
                priority = item.get("priority", "Medium")
                section = item.get("section", "Resume")
                st.markdown(f"### {index}. {section} — {priority} priority")
                st.write(f"**Issue:** {item.get('issue', '')}")
                st.write(f"**Recommendation:** {item.get('recommendation', '')}")
                if item.get("example"):
                    st.write(f"**Example/template:** {item.get('example')}")

        with st.expander("View extracted resume text"):
            st.text(resume_text)

    except Exception as exc:
        st.error("The analysis could not be completed.")
        st.exception(exc)


if __name__ == "__main__":
    main()
